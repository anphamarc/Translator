from pathlib import Path

import pytest

from translator_app.file_translate import UnsupportedFileError, output_path, translate_file


def test_txt(tmp_path, fake_engine):
    f = tmp_path / "a.txt"
    f.write_text("Hello\n\n123\nWorld\n", encoding="utf-8")
    out = translate_file(f, "en", "vi")
    assert out.name == "a_vi.txt"
    assert out.read_text(encoding="utf-8") == "[en>vi] Hello\n\n123\n[en>vi] World\n"


def test_markdown_keeps_markup_and_code(tmp_path, fake_engine):
    f = tmp_path / "a.md"
    f.write_text("# Title\n- item\n```\ncode here\n```\n", encoding="utf-8")
    out = translate_file(f, "en", "ko").read_text(encoding="utf-8")
    assert out == "# [en>ko] Title\n- [en>ko] item\n```\ncode here\n```\n"


def test_srt(tmp_path, fake_engine):
    f = tmp_path / "s.srt"
    f.write_text("1\n00:00:01,000 --> 00:00:02,000\nHello\nthere\n\n2\n00:00:03,000 --> 00:00:04,000\nBye\n")
    out = translate_file(f, "en", "vi").read_text(encoding="utf-8")
    assert out == (
        "1\n00:00:01,000 --> 00:00:02,000\n[en>vi] Hello there\n\n"
        "2\n00:00:03,000 --> 00:00:04,000\n[en>vi] Bye\n"
    )


def test_docx(tmp_path, fake_engine):
    import docx

    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Hello ").bold = True
    p.add_run("world")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Cell"
    d.sections[0].header.paragraphs[0].text = "Header"
    f = tmp_path / "d.docx"
    d.save(f)
    out = docx.Document(str(translate_file(f, "en", "ko")))
    assert out.paragraphs[0].text == "[en>ko] Hello world"
    assert out.paragraphs[0].runs[0].bold
    assert out.tables[0].cell(0, 0).text == "[en>ko] Cell"
    assert out.sections[0].header.paragraphs[0].text == "[en>ko] Header"


def test_pptx(tmp_path, fake_engine):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1)).table.cell(0, 0).text = "Cell"
    slide.notes_slide.notes_text_frame.text = "Note"
    f = tmp_path / "p.pptx"
    prs.save(f)
    out = Presentation(str(translate_file(f, "en", "vi")))
    texts = [s.text_frame.text for s in out.slides[0].shapes if s.has_text_frame]
    assert "[en>vi] Title" in texts
    table = [s for s in out.slides[0].shapes if s.has_table][0].table
    assert table.cell(0, 0).text == "[en>vi] Cell"
    assert out.slides[0].notes_slide.notes_text_frame.text == "[en>vi] Note"


def test_xlsx_skips_formulas(tmp_path, fake_engine):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = "Name", 5, "=A2*2"
    f = tmp_path / "x.xlsx"
    wb.save(f)
    ws = openpyxl.load_workbook(translate_file(f, "en", "vi")).active
    assert (ws["A1"].value, ws["A2"].value, ws["A3"].value) == ("[en>vi] Name", 5, "=A2*2")


def test_pdf_to_docx(tmp_path, fake_engine):
    import docx
    import pymupdf

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Hello PDF")
    f = tmp_path / "f.pdf"
    pdf.save(f)
    out = translate_file(f, "en", "vi")
    assert out.suffix == ".docx"
    assert "[en>vi] Hello PDF" in [p.text for p in docx.Document(str(out)).paragraphs]


def test_output_path_does_not_overwrite(tmp_path):
    src = tmp_path / "a.txt"
    (tmp_path / "a_vi.txt").write_text("x")
    assert output_path(src, "vi") == tmp_path / "a_vi (2).txt"


def test_unsupported(tmp_path, fake_engine):
    with pytest.raises(UnsupportedFileError):
        translate_file(tmp_path / "a.exe", "en", "vi")
