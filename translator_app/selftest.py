"""Tu kiem tra ban dong goi: dich that, OCR that, dich tep.

    Translator.exe --selftest   (ket qua ghi vao data/selftest.log)
"""

from __future__ import annotations

import tempfile
import traceback
from pathlib import Path

from translator_app import config


def run_selftest() -> int:
    config.prepare_environment()
    log_path = config.DATA_DIR / "selftest.log"
    lines: list[str] = []
    ok = True

    def check(name, fn):
        nonlocal ok
        try:
            lines.append(f"OK   {name}: {fn()}")
        except Exception:  # noqa: BLE001
            ok = False
            lines.append(f"FAIL {name}:\n{traceback.format_exc()}")

    from translator_app import models
    from translator_app.engine import get_engine

    check("models", lambda: f"thieu={models.missing_pairs()} tach_cau={models.sentence_models_installed()}"
          f" ocr={models.ocr_installed()}")
    eng = get_engine()
    check("en->vi", lambda: eng.translate_text("Hello, how are you today?", "en", "vi"))
    check("en->ko", lambda: eng.translate_text("Good morning, my friend.", "en", "ko"))
    check("ko->vi", lambda: eng.translate_text("오늘 날씨가 좋네요.", "ko", "vi"))
    check("vi->en", lambda: eng.translate_text("Xin chào, rất vui được gặp bạn.", "auto", "en"))

    def ocr_test():
        from PIL import Image, ImageDraw, ImageFont

        from translator_app import ocr

        img = Image.new("RGB", (700, 160), "white")
        ImageDraw.Draw(img).text((30, 40), "Hello world", fill="black", font=ImageFont.load_default(64))
        blocks, lang = ocr.recognize(img)
        text = ocr.blocks_to_text(blocks)
        assert "hello" in text.lower(), text
        return f"{lang}: {text}"

    check("ocr", ocr_test)

    def docx_test():
        import docx

        from translator_app.file_translate import translate_file

        tmp = Path(tempfile.mkdtemp())
        d = docx.Document()
        d.add_paragraph("This is a test document.")
        d.save(tmp / "t.docx")
        return docx.Document(str(translate_file(tmp / "t.docx", "en", "vi"))).paragraphs[0].text

    check("docx", docx_test)
    lines.append("PASS" if ok else "FAILED")
    log_path.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    return 0 if ok else 1
