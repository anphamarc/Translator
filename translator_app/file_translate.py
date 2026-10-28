"""Dich tep: txt, md, srt, docx, pptx, xlsx, pdf va anh.

Moi dinh dang duoc tach thanh cac "doan" (segment) kem ham ghi lai ban dich,
nho vay giu duoc dinh dang goc (font, bang, o tinh...) o muc toi da.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Callable

from translator_app.engine import CancelFn, ProgressFn, get_engine
from translator_app.languages import LANGUAGES

TEXT_EXT = {".txt", ".md", ".srt"}
OFFICE_EXT = {".docx", ".pptx", ".xlsx"}
PDF_EXT = {".pdf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
SUPPORTED_EXT = TEXT_EXT | OFFICE_EXT | PDF_EXT | IMAGE_EXT

Setter = Callable[[str], None]


class UnsupportedFileError(ValueError):
    pass


def output_path(src: Path, target: str, out_dir: Path | None = None, ext: str | None = None) -> Path:
    """vd. bao_cao.docx -> bao_cao_vi.docx (khong ghi de tep da co)."""
    folder = Path(out_dir) if out_dir else src.parent
    ext = ext or src.suffix
    candidate = folder / f"{src.stem}_{target}{ext}"
    n = 2
    while candidate.exists():
        candidate = folder / f"{src.stem}_{target} ({n}){ext}"
        n += 1
    return candidate


def _apply(segments: list[tuple[str, Setter]], src, dst, progress, cancelled) -> None:
    texts = [t for t, _ in segments]
    for (_, setter), result in zip(segments, get_engine().translate_batch(texts, src, dst, progress, cancelled)):
        setter(result)


# ---- van ban thuan ----------------------------------------------------------

def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp949", "cp1258", "cp1252"):
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
            continue
        return text
    return raw.decode("utf-8", errors="replace")


_MD_PREFIX = re.compile(r"^(\s*(?:#{1,6}\s+|[-*+]\s+(?:\[[ xX]\]\s+)?|\d+[.)]\s+|>\s*)*)(.*)$")


def _translate_plain(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    lines = read_text_file(path).splitlines()
    segments: list[tuple[str, Setter]] = []
    in_code = False
    is_md = path.suffix.lower() == ".md"
    for i, line in enumerate(lines):
        if is_md and line.lstrip().startswith("```"):
            in_code = not in_code
            continue
        if in_code or not line.strip():
            continue
        prefix, body = _MD_PREFIX.match(line).groups() if is_md else ("", line)

        def setter(t, i=i, prefix=prefix):
            lines[i] = prefix + t

        segments.append((body, setter))
    _apply(segments, src, dst, progress, cancelled)
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _translate_srt(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    blocks = re.split(r"\n\s*\n", read_text_file(path).replace("\r\n", "\n").strip())
    parsed = [b.split("\n") for b in blocks]
    segments: list[tuple[str, Setter]] = []
    for lines in parsed:
        # dong 1: so thu tu, dong 2: thoi gian, phan con lai: loi thoai
        start = 2 if len(lines) >= 2 and "-->" in lines[1] else 1 if "-->" in lines[0] else 0
        if len(lines) <= start:
            continue
        text = " ".join(lines[start:])

        def setter(t, lines=lines, start=start):
            lines[start:] = [t]

        segments.append((text, setter))
    _apply(segments, src, dst, progress, cancelled)
    out.write_text("\n\n".join("\n".join(lines) for lines in parsed) + "\n", encoding="utf-8")


# ---- Word / PowerPoint / Excel ---------------------------------------------

def _runs_segment(runs) -> tuple[str, Setter] | None:
    """Gop cac run cua mot doan; ban dich ghi vao run dau (giu dinh dang run dau)."""
    runs = list(runs)
    text = "".join(r.text for r in runs)
    if not runs or not text.strip():
        return None

    def setter(t):
        runs[0].text = t
        for r in runs[1:]:
            r.text = ""

    return text, setter


def _docx_paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _docx_paragraphs(cell)


def _translate_docx(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    import docx

    doc = docx.Document(str(path))
    containers = [doc]
    for section in doc.sections:
        # Chi lay dau/chan trang da co; truy cap phan chua co se tao moi no
        containers += [hf for hf in (section.header, section.footer) if not hf.is_linked_to_previous]
    segments, seen = [], set()
    for c in containers:
        for p in _docx_paragraphs(c):
            # o gop trong bang tra ve cung mot doan nhieu lan
            if id(p._p) in seen:
                continue
            seen.add(id(p._p))
            seg = _runs_segment(p.runs)
            if seg:
                segments.append(seg)
    _apply(segments, src, dst, progress, cancelled)
    doc.save(str(out))


def _pptx_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shapes(shape.shapes)
            continue
        if shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def _translate_pptx(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    segments = []
    for slide in prs.slides:
        paragraphs = list(_pptx_shapes(slide.shapes))
        if slide.has_notes_slide:
            paragraphs += list(slide.notes_slide.notes_text_frame.paragraphs)
        for p in paragraphs:
            seg = _runs_segment(p.runs)
            if seg:
                segments.append(seg)
    _apply(segments, src, dst, progress, cancelled)
    prs.save(str(out))


def _translate_xlsx(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(path))
    segments = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() and not cell.value.startswith("="):

                    def setter(t, cell=cell):
                        cell.value = t

                    segments.append((cell.value, setter))
    _apply(segments, src, dst, progress, cancelled)
    wb.save(str(out))


# ---- PDF va anh -------------------------------------------------------------

def _translate_pdf(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    """PDF -> Word (.docx). Trang scan (khong co chu) se duoc OCR."""
    import docx
    import pymupdf
    from PIL import Image

    pdf = pymupdf.open(str(path))
    pages: list[list[str]] = []
    for page in pdf:
        blocks = [b[4] for b in page.get_text("blocks", sort=True) if b[6] == 0]
        blocks = [re.sub(r"\s*\n\s*", " ", b).strip() for b in blocks]
        blocks = [b for b in blocks if b]
        if not blocks:
            from translator_app import ocr

            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            blocks = [b.text for b in ocr.recognize(img, src)[0]]
        pages.append(blocks)
        if cancelled and cancelled():
            from translator_app.engine import Cancelled

            raise Cancelled()

    segments = []
    for blocks in pages:
        for i in range(len(blocks)):

            def setter(t, blocks=blocks, i=i):
                blocks[i] = t

            segments.append((blocks[i], setter))
    _apply(segments, src, dst, progress, cancelled)

    doc = docx.Document()
    for n, blocks in enumerate(pages, 1):
        if n > 1:
            doc.add_page_break()
        doc.add_heading(f"Trang {n}", level=2)
        for b in blocks:
            doc.add_paragraph(b)
    doc.save(str(out))


def _translate_image(path: Path, out: Path, src, dst, progress, cancelled) -> None:
    from PIL import Image

    from translator_app import ocr

    img = Image.open(path)
    blocks, detected = ocr.recognize(img, src)
    segments = []
    for b in blocks:

        def setter(t, b=b):
            b.translation = t

        segments.append((b.text, setter))
    _apply(segments, detected, dst, progress, cancelled)
    ocr.render_translation(img, blocks, dst).save(out)


# ---- diem vao ---------------------------------------------------------------

def translate_file(
    path: str | Path,
    src: str,
    dst: str,
    out_dir: str | Path | None = None,
    progress: ProgressFn | None = None,
    cancelled: CancelFn | None = None,
) -> Path:
    """Dich mot tep va tra ve duong dan tep ket qua."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXT:
        raise UnsupportedFileError(f"Chưa hỗ trợ định dạng {ext or '(không có đuôi)'}")
    if dst not in LANGUAGES:
        raise ValueError(f"Ngôn ngữ đích không hợp lệ: {dst}")
    if ext in TEXT_EXT:
        handler = _translate_srt if ext == ".srt" else _translate_plain
        out = output_path(path, dst, out_dir)
    elif ext in PDF_EXT:
        handler, out = _translate_pdf, output_path(path, dst, out_dir, ".docx")
    elif ext in IMAGE_EXT:
        handler, out = _translate_image, output_path(path, dst, out_dir, ".png")
    else:
        handler = {".docx": _translate_docx, ".pptx": _translate_pptx, ".xlsx": _translate_xlsx}[ext]
        out = output_path(path, dst, out_dir)
    handler(path, out, src, dst, progress, cancelled)
    return out
